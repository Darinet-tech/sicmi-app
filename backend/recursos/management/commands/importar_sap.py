import unicodedata
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from accounts.models import UnidadOrganizativa
from recursos.models import Almacen, Material, Existencia

# Nombres posibles del encabezado en el Excel (sin tildes y en minúsculas).
# Si su archivo usa otro nombre, agréguelo a la lista correspondiente.
CANDIDATOS = {
    'material': ['material', 'codigo sap', 'numero de material', 'codigo material'],
    'descripcion': ['texto breve de material', 'texto breve', 'descripcion', 'denominacion'],
    'almacen': ['almacen', 'alm.'],
    'unidad': ['um base', 'unidad de medida base', 'um', 'unidad medida'],
    'libre': ['libre utilizacion', 'libre utiliz.', 'stock libre utilizacion'],
    'valor': ['valor libre util.', 'valor libre utilizacion', 'valor libre utiliz.'],
}


def normalizar(texto):
    t = unicodedata.normalize('NFD', str(texto or '')).encode('ascii', 'ignore').decode()
    return ' '.join(t.lower().split())


def a_decimal(valor):
    """Convierte números de Excel o texto SAP ('1.234,50', '5-') a Decimal."""
    if valor is None or valor == '':
        return Decimal('0')
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    s = str(valor).strip().replace(' ', '')
    if s.endswith('-'):  # SAP escribe los negativos con el signo al final
        s = '-' + s[:-1]
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal('0')


def a_texto(valor):
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return '' if valor is None else str(valor).strip()


class Command(BaseCommand):
    help = 'Importa existencias en libre utilización desde un Excel exportado de SAP.'

    def add_arguments(self, parser):
        parser.add_argument('archivo', help='Ruta del archivo .xlsx')
        parser.add_argument('--uo', required=True, help='Acrónimo de la UO (ej. DTIJ)')
        parser.add_argument('--almacen', help='Código de almacén, si el archivo no trae esa columna')

    def handle(self, *args, **opts):
        try:
            uo = UnidadOrganizativa.objects.get(acronimo=opts['uo'])
        except UnidadOrganizativa.DoesNotExist:
            raise CommandError(f"No existe la UO con acrónimo '{opts['uo']}'.")

        hoja = load_workbook(opts['archivo'], read_only=True, data_only=True).active
        filas = hoja.iter_rows(values_only=True)
        encabezado = next(filas)
        normales = [normalizar(h) for h in encabezado]

        cols = {}
        for campo, candidatos in CANDIDATOS.items():
            for i, h in enumerate(normales):
                if h in candidatos:
                    cols[campo] = i
                    break

        faltan = [c for c in ('material', 'descripcion', 'libre') if c not in cols]
        if faltan:
            raise CommandError(f'No encontré las columnas {faltan}. Encabezados del archivo: {list(encabezado)}')
        if 'almacen' not in cols and not opts['almacen']:
            raise CommandError('El archivo no tiene columna de almacén: use --almacen CODIGO.')

        def dato(fila, campo):
            i = cols.get(campo)
            return fila[i] if i is not None and i < len(fila) else None

        leidas = importadas = omitidas = 0
        vistos = set()
        almacenes = {}

        with transaction.atomic():
            for fila in filas:
                codigo = a_texto(dato(fila, 'material'))
                if not codigo:
                    continue
                leidas += 1

                cantidad = a_decimal(dato(fila, 'libre'))
                if cantidad <= 0:  # solo libre utilización con existencia
                    omitidas += 1
                    continue

                cod_alm = a_texto(dato(fila, 'almacen')) or opts['almacen']
                if cod_alm not in almacenes:
                    almacenes[cod_alm], _ = Almacen.objects.get_or_create(uo=uo, codigo=cod_alm)
                almacen = almacenes[cod_alm]

                material, _ = Material.objects.update_or_create(
                    codigo_sap=codigo,
                    defaults={
                        'descripcion': a_texto(dato(fila, 'descripcion'))[:255],
                        'unidad_medida': a_texto(dato(fila, 'unidad'))[:15],
                    },
                )

                valor = a_decimal(dato(fila, 'valor')) if 'valor' in cols else Decimal('0')
                unitario = (valor / cantidad).quantize(Decimal('0.0001')) if valor > 0 else None

                Existencia.objects.update_or_create(
                    almacen=almacen, material=material,
                    defaults={'cantidad_libre': cantidad, 'valor_unitario': unitario},
                )
                vistos.add((almacen.id, material.id))
                importadas += 1

            # Sincronizar: borrar lo que ya no está en el archivo
            eliminadas = 0
            for almacen in almacenes.values():
                for ex in Existencia.objects.filter(almacen=almacen):
                    if (ex.almacen_id, ex.material_id) not in vistos:
                        ex.delete()
                        eliminadas += 1

        self.stdout.write(self.style.SUCCESS(
            f'Listo. Filas leídas: {leidas} | importadas: {importadas} | '
            f'sin libre utilización: {omitidas} | existencias eliminadas: {eliminadas}'
        ))